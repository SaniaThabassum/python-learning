import sys 
from openpyxl import Workbook 
def create_multiplication_table(n): 
    wb = Workbook() 
    ws = wb.active 
    ws.title = f"{n}x{n} Table" 
# Create table headers 
    for i in range(1, n + 1): 
        ws.cell(row=1, column=i + 1, value=i) 
        ws.cell(row=i + 1, column=1, value=i) 
# Fill multiplication table 
    for row in range(1, n + 1): 
        for col in range(1, n + 1): 
            ws.cell(row=row + 1, column=col + 1, value=row * col) 
# Save file 
    file_name = "multiplication_table.xlsx" 
    wb.save(file_name) 
    print(f"Multiplication table saved as '{file_name}'") 
 
try: 
    N = int(input("enter value of N")) 
    if N <= 0: 
 
  
        print("Please enter a positive integer") 
    else: 
        create_multiplication_table(N) 
except ValueError: 
    print("Please enter a valid integer") 